#!/usr/bin/env python3
"""Convert FC APS Excel workbooks into a standard ccAPPS dataset JSON.

ccAPPS mapping
==============
* 设备 / 模具 / 人员 are three resource categories. Personnel are one pool
  per 工段; size follows the shift calendar rather than a single maximum.
* 标准工时.CT时间 is cycle time in seconds (matches 小时产量 ≈ 3600/CT).
  That value is ``duration_per``. ``duration`` is setup time and stays empty.
* 物料机台对照表 separators: ``/`` (and comma) is a list of interchangeable
  machines and becomes alternate children. ``&`` is one inseparable group:
  one child that loads every member machine together (not pick-one).
* 「N工段装配线」 is a work-center pool, not an occupy-all group: no station
  list exists, so the line resource is infinite and parallel jobs are limited
  only by that 工段's PERSON calendar and 标准工时.配置人数.
* Every time_per leaf loads 设备 (only from 物料机台对照表) + 人员 (by 工段)
  + 模具. No neighbor/mold machine inference. 工段 comes from the loaded
  machines, else from 未入库.生产车间 when present.
* Alternate parents must not consume/produce material (ccAPPS deprecation).
  End-item and BOM start flows go on each time_per child; the parent only
  selects among children.
* Personnel calendars follow 人数.xlsx dates only (here 9/10 08:00 →
  9/13 08:00). Outside those windows headcount is 0. Night headcount 0
  means that window is not used. ``sizemaximum`` = floor(batch seconds / CT):
  制造部5工段 is 1 hour, every other 工段 is 2 hours. If an operation
  touches 5工段, the 1-hour cap wins. Economic lot is NOT ``sizemultiple``
  — remainders must still schedule later.
* 未入库.未入库数量 = still unfinished qty that must be re-planned (open
  demand, priority 1; due = 计划结束日期 + 1 day at 00:00). It is NOT
  onhand. Workshop hint still comes from 生产车间.
* Order-delivery products keep only real warehouse ATP from 库存.xlsx:
  max(0, 即时库存 + 已完工未入库 - 未发货). 未发货 is reserved for ship
  orders and is not free stock. BOM components that are NOT order items
  still get large assumed onhand (原材料 / 自制中间料充足). Pure BOM-make
  items that are not order items stay buffer-only (no routing).
* 物料模具对照表: ``1*N`` same-cavity multi-shot uses 单个CT as duration_per;
  ``1+1`` / same 模具+分组编号 co-produce on one operation (equal qty);
  same 配套编号 also co-produce (molds may differ, all required).
* 产品外购 / 外购 items are supply only (buffer, no routing).
"""
from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).parent
SOURCE = "fc_aps"
LOCATION = "FC_FACTORY"
MAKE_ATTRS = {"产品自制", "自制", "自制件"}
# Temporary planning assumption: BOM-exploded make items are covered by stock.
ASSUMED_MAKE_STOCK = 1_000_000
# 对照表里的「N工段装配线」是工段工作中心（资源池），不是机台.xlsx 里的单台，
# 也不是 B03&B04 那种整组锁定。无限产能，只靠该工段人员池限制并行。
ASSEMBLY_LINES = {
    "1工段装配线": "制造部1工段",
    "2工段装配线": "制造部2工段",
    "3工段装配线": "制造部3工段",
    "4工段装配线": "制造部4工段",
    "5工段装配线": "制造部5工段",
}
# Only ``/`` and commas are alternatives. ``&`` stays one group token.
MACHINE_SPLIT = re.compile(r"[,，/]+")
OP_SAFE = re.compile(r"[^0-9A-Za-z_.+-]+")


def rows(name):
    ws = load_workbook(ROOT / name, read_only=True, data_only=True).active
    return list(ws.iter_rows(values_only=True))[1:]


def num(x, default=0.0):
    try:
        if x is None or x == "":
            return float(default)
        return float(x)
    except (TypeError, ValueError):
        return float(default)


def text(x):
    if x is None:
        return ""
    return str(x).strip()


def parse_cavity(raw):
    """Return (kind, pieces) where kind is 'same' (*), 'diff' (+), or 'one'."""
    s = text(raw)
    if not s:
        return "one", 1
    if "*" in s:
        parts = [p.strip() for p in s.split("*") if p.strip()]
        try:
            pieces = 1
            for p in parts:
                pieces *= max(int(float(p)), 1)
            return "same", pieces
        except (TypeError, ValueError):
            return "same", 1
    if "+" in s:
        parts = [p.strip() for p in s.split("+") if p.strip()]
        try:
            return "diff", sum(max(int(float(p)), 1) for p in parts)
        except (TypeError, ValueError):
            return "diff", 2
    return "one", 1


def op_safe(token):
    return OP_SAFE.sub("_", token).strip("_") or "X"


def iso_datetime(x):
    if not x:
        return None
    if isinstance(x, datetime):
        return x.isoformat()
    raw = str(x).strip()
    for fmt in ("%Y/%m/%d %H:%M:%S", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, fmt).isoformat()
        except ValueError:
            continue
    return raw


def iso_date_only(x, days_offset=0):
    """Plan due as calendar date at 00:00:00 (drop time-of-day)."""
    if not x:
        return None
    if isinstance(x, datetime):
        d = x.date()
    else:
        raw = str(x).strip()
        d = None
        for fmt in ("%Y/%m/%d %H:%M:%S", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
            try:
                d = datetime.strptime(raw, fmt).date()
                break
            except ValueError:
                continue
        if d is None:
            return None
    if days_offset:
        d = d + timedelta(days=days_offset)
    return d.isoformat() + "T00:00:00"


def split_machines(raw):
    if not raw:
        return []
    return [part.strip() for part in MACHINE_SPLIT.split(str(raw)) if part.strip()]


def machine_members(token):
    return [part.strip() for part in str(token).split("&") if part.strip()]


def op_suffix(token):
    return token.replace("&", "+")


def trailing_num(code):
    tail = code.split(".")[-1]
    digits = "".join(ch for ch in tail if ch.isdigit())
    return int(digits) if digits else 0


def parse_day(value):
    """Normalize 人数.xlsx dates like 2026/9/10 → date."""
    if isinstance(value, datetime):
        return value.date()
    text_v = str(value or "").strip().split()[0]
    for fmt in ("%Y/%m/%d", "%Y-%m-%d"):
        try:
            return datetime.strptime(text_v, fmt).date()
        except ValueError:
            continue
    raise ValueError("bad 人数 date: %r" % (value,))


def clock_seconds(t):
    parts = str(t or "00:00").strip().split(":")
    h = int(float(parts[0]))
    m = int(float(parts[1])) if len(parts) > 1 else 0
    s = int(float(parts[2])) if len(parts) > 2 else 0
    return h * 3600 + m * 60 + s


def normalize_clock(t, default="00:00:00"):
    text_v = str(t or default).strip()
    if len(text_v) == 5:
        text_v += ":00"
    return text_v


def overnight_windows(day, start, end):
    """Split a shift into dated ccAPPS windows (overnight ends next morning)."""
    start = normalize_clock(start, "00:00:00")
    end = normalize_clock(end, "23:59:59")
    if start < end:
        return [(day, start, end)]
    nxt = day + timedelta(days=1)
    return [(day, start, "23:59:59"), (nxt, "00:00:00", end)]


def add_bucket(buckets, calendar, value, day, start, end):
    # Unique key is (calendar, startdate, enddate, priority). Day/night and
    # overnight splits must not share priority 0.
    used = {b["priority"] for b in buckets if b["calendar"] == calendar}
    priority = (max(used) + 1) if used else 0
    for when, starttime, endtime in overnight_windows(day, start, end):
        day_s = when.isoformat()
        buckets.append(
            {
                "calendar": calendar,
                "value": value,
                "startdate": day_s,
                "enddate": day_s,
                "starttime": starttime,
                "endtime": endtime,
                "priority": priority,
                "monday": True,
                "tuesday": True,
                "wednesday": True,
                "thursday": True,
                "friday": True,
                "saturday": True,
                "sunday": True,
                "source": SOURCE,
            }
        )
        priority += 1


def op_resource(operation, resource, quantity, name, priority):
    return {
        "operation": operation,
        "resource": resource,
        "quantity": quantity,
        "name": name,
        "priority": priority,
        "search": "PRIORITY",
        "source": SOURCE,
    }


def build():
    item_name = {}
    item_category = {}
    make_codes = set()
    for code, org, name, lifecycle, attr, econ, *_ in rows("物料.xlsx"):
        code = text(code)
        if not code:
            continue
        item_name[code] = text(name) or code
        attr = text(attr)
        if attr in MAKE_ATTRS:
            make_codes.add(code)
            item_category[code] = "自制"
        else:
            item_category.setdefault(code, attr or "外购")

machines = {}
    for name, model, workshop in rows("机台.xlsx"):
        name = text(name)
        if name:
            machines[name] = (text(workshop), text(model) or name)
    for line, workshop in ASSEMBLY_LINES.items():
        machines.setdefault(line, (workshop, line))

molds = {}
    for mold, name, cavities, quantity in rows("模具.xlsx"):
        mold = text(mold)
        if mold:
            molds[mold] = (text(name) or mold, max(num(quantity, 1), 1))

    # Exact 人数 rows: plan only inside these dated windows (9/10 08:00→9/13 08:00).
    shift_rows = []
    for day, workshop, shift, start, end, count in rows("人数.xlsx"):
        workshop = text(workshop)
        shift = text(shift) or "白班"
        if not workshop:
            continue
        shift_rows.append(
            (
                parse_day(day),
                workshop,
                shift,
                start or "08:00",
                end or "20:00",
                num(count, 0),
            )
        )

    workshops = sorted({ws for _, ws, *_ in shift_rows})

    # One MO is at most 1h on 5工段, 2h everywhere else.
    BATCH_SEC = {
        "制造部5工段": 3600,
    }
    DEFAULT_BATCH_SEC = 7200
    # Filled after mold_info so 1*N can scale shot CT → piece CT.
    duration_per = {}
staffing = {}
    std_hours = {}
    for code, ct, output, people, cavities, single_ct, *_ in rows("标准工时.xlsx"):
        code = text(code)
        if not code:
            continue
        std_hours[code] = (num(ct, 0), num(single_ct, 0), num(output, 0), num(people, 1))
        staffing[code] = num(people, 1)

    machine_priority = {}
    for code, rated, spare, primary in rows("物料机台对照表.xlsx"):
        code = text(code)
        if not code:
            continue
        ranked = machine_priority.setdefault(code, {})
        for raw, prio in ((rated, 1), (spare, 2)):
            for token in split_machines(raw):
                ranked[token] = min(ranked.get(token, prio), prio)
        # 主机台 is the preferred machine inside a group, not a new split.
        for token in split_machines(primary):
            members = machine_members(token)
            covered = False
            for existing, prio in list(ranked.items()):
                existing_members = machine_members(existing)
                if token == existing or (members and all(m in existing_members for m in members)):
                    ranked[existing] = min(prio, 1)
                    covered = True
            if not covered:
                ranked[token] = min(ranked.get(token, 1), 1)

    # code -> {mold, cavity, group, kit, kind, pieces}
    mold_info = {}
mold_map = {}
    for code, mold, cavity, group, kit, *_ in rows("物料模具对照表.xlsx"):
        code, mold = text(code), text(mold)
        if not code or not mold:
            continue
        kind, pieces = parse_cavity(cavity)
        mold_info[code] = {
            "mold": mold,
            "cavity": text(cavity),
            "group": text(group),
            "kit": text(kit),
            "kind": kind,
            "pieces": pieces,
        }
        mold_map.setdefault(code, [])
        if mold not in mold_map[code]:
            mold_map[code].append(mold)
        molds.setdefault(mold, (mold, 1))

    for code, (shot_ct, piece_ct, out, _people) in std_hours.items():
        pieces = 1
        info = mold_info.get(code)
        if info and info["kind"] == "same":
            pieces = max(info["pieces"], 1)
        if piece_ct > 0:
            duration_per[code] = piece_ct
        elif shot_ct > 0 and pieces > 1:
            duration_per[code] = shot_ct / pieces
        elif out > 0:
            duration_per[code] = 3600.0 / out
        else:
            duration_per[code] = shot_ct

    # Explicit 工段 hint when 机台对照 is missing (人员按工段挂，不靠邻项推机台).
    # 未入库数量 = unfinished qty to re-plan (not onhand).
    item_workshop_votes = defaultdict(Counter)
    unin_rows = []
stock = {}
    stock_imm = defaultdict(float)
    stock_fin = defaultdict(float)
    stock_uns = defaultdict(float)
    for code, org, immediate, finished, unshipped in rows("库存.xlsx"):
        code = text(code)
        if not code:
            continue
        stock_imm[code] += num(immediate)
        stock_fin[code] += num(finished)
        stock_uns[code] += num(unshipped)
    for code in set(stock_imm) | set(stock_fin) | set(stock_uns):
        # ATP: on-hand + finished-not-receipted, minus reserved-to-ship.
        stock[code] = max(stock_imm[code] + stock_fin[code] - stock_uns[code], 0.0)
    for row in rows("未入库.xlsx"):
        code, name, qty, org, plan_qty, start, end, workshop, machine, order, *_ = list(row) + [None] * 10
        code = text(code)
        qty = num(qty)
        workshop = text(workshop)
        if code and workshop:
            item_workshop_votes[code][workshop] += 1
        if code and qty > 0:
            unin_rows.append(
                {
                    "code": code,
                    "name": text(name) or code,
                    "qty": qty,
                    "org": text(org),
                    "plan_qty": num(plan_qty),
                    "start": start,
                    "end": end,
                    "workshop": workshop,
                    "machine": text(machine),
                    "order": text(order),
                }
            )
            item_name.setdefault(code, text(name) or code)
    item_workshop = {
        code: votes.most_common(1)[0][0]
        for code, votes in item_workshop_votes.items()
        if votes
    }

    bom_rows = []
    for parent, pname, child, cname, numerator, denominator, org, *_ in rows("BOM.xlsx"):
        parent, child = text(parent), text(child)
        if not parent or not child:
            continue
        bom_rows.append((parent, child, num(numerator, 1), max(num(denominator, 1), 1e-12), text(pname), text(cname)))
        item_name.setdefault(parent, pname or parent)
        item_name.setdefault(child, cname or child)
        if parent not in item_category:
            item_category[parent] = "自制"
        if child not in item_category:
            item_category[child] = "外购"

    bom_children = defaultdict(list)
    for parent, child, *_ in bom_rows:
        bom_children[parent].append(child)

    def is_make(code):
        if code in make_codes:
            return True
        return item_category.get(code) in MAKE_ATTRS or item_category.get(code) == "自制"

    # Open alarm + unfinished 未入库 自制 → build routing. BOM-descendant 自制 → assume stock covers.
    order_items = set()
    for code, org, ship, qty in rows("报警表.xlsx"):
        code = text(code)
        if code and num(qty) > 0:
            order_items.add(code)
            item_name.setdefault(code, code)
    for row in unin_rows:
        code = row["code"]
        order_items.add(code)
        item_name.setdefault(code, row["name"])

    produce = {code for code in order_items if is_make(code)}

    # Walk full BOM under produce. Pure BOM-make (not order items) → no routing.
    # Assume raw / intermediate supply is sufficient, but order-delivery products
    # keep only real warehouse stock (do not inflate).
    bom_make = set()
    bom_components = set()
    stack = list(produce)
    visited = set()
    while stack:
        code = stack.pop()
        if code in visited:
            continue
        visited.add(code)
        for child in bom_children.get(code, []):
            bom_components.add(child)
            if child not in visited:
                stack.append(child)
            if is_make(child) and child not in produce:
                bom_make.add(child)

    for code in bom_make:
        item_name.setdefault(code, code)
        item_category.setdefault(code, "自制")

    for code in bom_components:
        if code in order_items:
            continue
        stock[code] = max(stock.get(code, 0), ASSUMED_MAKE_STOCK)

    # Sync groups: same 模具+分组编号, or same 配套编号 (molds may differ).
    group_members = defaultdict(list)
    kit_members = defaultdict(list)
    for code, info in mold_info.items():
        if info["group"]:
            group_members[(info["mold"], info["group"])].append(code)
        if info["kit"]:
            kit_members[info["kit"]].append(code)

    sync_groups = []  # list of {id, members, molds, primary}
    seen_sync = set()
    for (mold, group), members in sorted(group_members.items()):
        members = sorted(set(members))
        if len(members) < 2:
            continue
        if not (set(members) & produce):
            continue
        gid = "G_%s_%s" % (op_safe(mold), op_safe(group))
        if gid in seen_sync:
            continue
        seen_sync.add(gid)
        sync_groups.append(
            {
                "id": gid,
                "members": members,
                "molds": [mold],
                "primary": sorted(set(members) & produce)[0],
            }
        )
    for kit, members in sorted(kit_members.items()):
        members = sorted(set(members))
        if len(members) < 2:
            continue
        if not (set(members) & produce):
            continue
        gid = "K_%s" % op_safe(kit)
        if gid in seen_sync:
            continue
        seen_sync.add(gid)
        molds_in = []
        for code in members:
            m = mold_info[code]["mold"]
            if m not in molds_in:
                molds_in.append(m)
        sync_groups.append(
            {
                "id": gid,
                "members": members,
                "molds": molds_in,
                "primary": sorted(set(members) & produce)[0],
            }
        )

    sync_of = {}
    for grp in sync_groups:
        for code in grp["members"]:
            sync_of[code] = grp
            item_name.setdefault(code, code)
            item_category.setdefault(code, "自制")
            # Co-product partners need to exist even if not open orders.
            if code not in produce and is_make(code):
                # Keep as buffer consumer/producer via shared op; do not add
                # independent routing. Ensure item/buffer exist.
                pass

    item_codes = set(item_name) | set(stock) | {p for p, *_ in bom_rows} | {c for _, c, *_ in bom_rows} | produce | order_items | bom_make
    for grp in sync_groups:
        item_codes.update(grp["members"])
    for code in produce:
        item_name.setdefault(code, code)
        item_category.setdefault(code, "自制")

    calendars = [
        {"name": "FC_24X7", "description": "默认可用", "defaultvalue": 100, "source": SOURCE},
    ]
    buckets = [
        {
            "calendar": "FC_24X7",
            "value": 100,
            "startdate": "2024-01-01",
            "enddate": "2035-12-31",
            "starttime": "00:00:00",
            "endtime": "23:59:59",
            "source": SOURCE,
        }
    ]
    for workshop in workshops:
        person_cal = "CAL_PERSON_" + workshop
        shift_cal = "CAL_SHIFT_" + workshop
        calendars.append({"name": person_cal, "description": workshop + "出勤人数", "defaultvalue": 0, "source": SOURCE})
        calendars.append({"name": shift_cal, "description": workshop + "班次", "defaultvalue": 0, "source": SOURCE})
        for day, ws, _shift, start, end, count in shift_rows:
            if ws != workshop:
                continue
            add_bucket(buckets, person_cal, count, day, start, end)
            if count > 0:
                add_bucket(buckets, shift_cal, 100, day, start, end)

    resources = []
    for machine, (workshop, model) in machines.items():
        row = {
            "name": machine,
            "description": ("%s %s" % (workshop, model)).strip(),
            "category": "设备",
            "type": "default",
            "maximum": 1,
            "efficiency": 100,
            "constrained": True,
            "location": LOCATION,
            "source": SOURCE,
        }
        if workshop in workshops:
            row["available"] = "CAL_SHIFT_" + workshop
        if machine in ASSEMBLY_LINES:
            row["type"] = "infinite"
            row["constrained"] = False
            row["subcategory"] = "装配工作中心"
            row["description"] = workshop + "装配工作中心（资源池，不整组锁定；并行靠人员）"
        resources.append(row)
    for mold, (name, quantity) in molds.items():
        resources.append(
            {
                "name": "MOLD_" + mold,
                "description": name,
                "category": "模具",
                "type": "default",
                "maximum": quantity,
                "efficiency": 100,
                "constrained": True,
                "location": LOCATION,
                "source": SOURCE,
            }
        )
    for workshop in workshops:
        peak = max(count for _day, ws, _shift, _s, _e, count in shift_rows if ws == workshop)
        resources.append(
            {
                "name": "PERSON_" + workshop,
                "description": workshop + "人员",
                "category": "人员",
                "type": "default",
                "maximum": peak,
                "maximum_calendar": "CAL_PERSON_" + workshop,
                "available": "CAL_SHIFT_" + workshop,
                "efficiency": 100,
                "constrained": True,
                "location": LOCATION,
                "source": SOURCE,
            }
        )

    operations = []
    operation_resources = []
    operation_materials = []

    def attach_molds(op_name, code):
        molds_for = mold_map.get(code, [])
        if len(molds_for) <= 1:
            for mold in molds_for:
                operation_resources.append(op_resource(op_name, "MOLD_" + mold, 1, "模具", 1))
            return
        for mold in molds_for:
            operation_resources.append(op_resource(op_name, "MOLD_" + mold, 1, "模具_" + op_safe(mold), 1))

    def attach_molds_list(op_name, mold_list):
        mold_list = [m for m in mold_list if m]
        if len(mold_list) <= 1:
            for mold in mold_list:
                operation_resources.append(op_resource(op_name, "MOLD_" + mold, 1, "模具", 1))
            return
        for mold in mold_list:
            operation_resources.append(op_resource(op_name, "MOLD_" + mold, 1, "模具_" + op_safe(mold), 1))

    def known_members(token):
        return [m for m in machine_members(token) if m in machines]

    def attach_people(op_name, code, workshop):
        if workshop in workshops:
            operation_resources.append(
                op_resource(op_name, "PERSON_" + workshop, staffing.get(code, 1), "人员", 1)
            )

    def attach_people_list(op_name, code, workshop_list):
        seen = []
        for workshop in workshop_list:
            if workshop in workshops and workshop not in seen:
                seen.append(workshop)
        if len(seen) <= 1:
            if seen:
                attach_people(op_name, code, seen[0])
            return
        for workshop in seen:
            operation_resources.append(
                op_resource(op_name, "PERSON_" + workshop, staffing.get(code, 1), "人员_" + workshop, 1)
            )

    def attach_equipment(op_name, code, token):
        members = known_members(token)
        if not members:
            return
        if len(members) == 1:
            operation_resources.append(op_resource(op_name, members[0], 1, "设备", 1))
            attach_people(op_name, code, machines[members[0]][0])
            return
        # ``&`` group: unique names so every machine is required, not alternate.
        for machine in members:
            operation_resources.append(op_resource(op_name, machine, 1, "设备_" + machine, 1))
        attach_people_list(op_name, code, [machines[m][0] for m in members])

    def workshops_for_token(token):
        return {machines[m][0] for m in known_members(token)}

    def time_fields(code, workshops_used=None):
        ct = duration_per.get(code)
        row = {"duration": None, "duration_per": ct if ct and ct > 0 else None}
        # Do not set sizemultiple from 经济批量: leftover qty smaller than the
        # multiple cannot form an MO and stays unmet even when capacity later.
        # 5工段: one MO ≤ 1 hour. Other 工段: one MO ≤ 2 hours.
        if ct and ct > 0:
            if workshops_used:
                batch_sec = min(BATCH_SEC.get(w, DEFAULT_BATCH_SEC) for w in workshops_used)
            else:
                batch_sec = DEFAULT_BATCH_SEC
            row["sizemaximum"] = max(1, int(batch_sec // ct))
        return row

    def attach_people_for_code(op_name, code, workshops_used=None):
        """人员按工段：优先用机台所属工段，否则用未入库.生产车间。"""
        if workshops_used:
            attach_people_list(op_name, code, list(workshops_used))
            return
        hint = item_workshop.get(code)
        if hint:
            attach_people(op_name, code, hint)

    def add_end_materials(op_name, codes):
        for code in codes:
            operation_materials.append(
                {
                    "operation": op_name,
                    "item": code,
                    "location": LOCATION,
                    "quantity": 1,
                    "type": "end",
                    "priority": 1,
                    "source": SOURCE,
                }
            )

    def build_routing_for(codes, op_base, description, mold_list, primary):
        """Build time_per / alternate routing producing all codes as co-products.

        Returns leaf time_per names that carry end/BOM material flows.
        Alternate parents only select among children and must not hold materials.
        """
        ranked = machine_priority.get(primary, {})
        # Prefer a member that actually has machine map.
        for code in codes:
            if machine_priority.get(code):
                ranked = machine_priority[code]
                primary = code
                break
        candidates = [m for m in ranked if known_members(m)]
        candidates.sort(key=lambda m: (ranked[m], m))
        staff_code = max(codes, key=lambda c: staffing.get(c, 1))
        # Shot/piece time: use max duration among members (conservative).
        cts = [duration_per[c] for c in codes if duration_per.get(c)]
        if cts:
            duration_per[primary] = max(cts)

        if not candidates:
            no_equipment.append(primary)
            extras = time_fields(primary, {item_workshop[primary]} if item_workshop.get(primary) else set())
            operations.append(
                {
                    "name": op_base,
                    "type": "time_per",
                    "description": description,
                    "item": primary,
                    "location": LOCATION,
                    "priority": 1,
                    "sizeminimum": 1,
                    "source": SOURCE,
                    **extras,
                }
            )
            attach_people_for_code(op_base, staff_code)
            attach_molds_list(op_base, mold_list)
            add_end_materials(op_base, codes)
            return [op_base]

        if len(candidates) <= 1:
            token = candidates[0]
            workshops_used = workshops_for_token(token)
            extras = time_fields(primary, workshops_used)
            operations.append(
                {
                    "name": op_base,
                    "type": "time_per",
                    "description": description,
                    "item": primary,
                    "location": LOCATION,
                    "priority": 1,
                    "sizeminimum": 1,
                    "source": SOURCE,
                    **extras,
                }
            )
            attach_equipment(op_base, staff_code, token)
            attach_molds_list(op_base, mold_list)
            add_end_materials(op_base, codes)
            return [op_base]

        operations.append(
            {
                "name": op_base,
                "type": "alternate",
                "description": description,
                "item": primary,
                "location": LOCATION,
                "priority": 1,
                "search": "PRIORITY",
                "sizeminimum": 1,
                "source": SOURCE,
            }
        )
        leaves = []
        for token in candidates:
            child = op_base + "__" + op_suffix(token)
            extras = time_fields(primary, workshops_for_token(token))
            operations.append(
                {
                    "name": child,
                    "type": "time_per",
                    "description": "%s @ %s" % (description, token),
                    "item": None,
                    "location": LOCATION,
                    "owner": op_base,
                    "priority": ranked[token],
                    "sizeminimum": 1,
                    "source": SOURCE,
                    **extras,
                }
            )
            attach_equipment(child, staff_code, token)
            attach_molds_list(child, mold_list)
            add_end_materials(child, codes)
            leaves.append(child)
        return leaves

    no_equipment = []
    producing_ops = {}
    producing_leaves = {}
    handled = set()

    # 1) Co-produce sync groups (分组 / 配套).
    for grp in sync_groups:
        members = grp["members"]
        primary = grp["primary"]
        op_base = "FC_OP_SYNC_" + grp["id"]
        desc = "联产:" + "+".join(item_name.get(c, c) for c in members[:4])
        if len(members) > 4:
            desc += "…"
        leaves = build_routing_for(members, op_base, desc, grp["molds"], primary)
        for code in members:
            producing_ops[code] = op_base
            producing_leaves[code] = leaves
            handled.add(code)

    # 2) Remaining open-order make items (solo routings).
    for code in sorted(produce):
        if code in handled:
            continue
        name = item_name.get(code, code)
        op_base = "FC_OP_" + code
        leaves = build_routing_for([code], op_base, name, mold_map.get(code, []), code)
        producing_ops[code] = op_base
        producing_leaves[code] = leaves

    for parent, child, numerator, denominator, *_ in bom_rows:
        if parent not in producing_leaves:
            continue
        qty = -numerator / denominator
        for leaf in producing_leaves[parent]:
            operation_materials.append(
                {
                    "operation": leaf,
                    "item": child,
                    "location": LOCATION,
                    "quantity": qty,
                    "type": "start",
                    "priority": 1,
                    "source": SOURCE,
                }
            )

    buffers = []
    for code in sorted(item_codes):
        buffers.append(
            {
                "item": code,
                "location": LOCATION,
                "batch": "",
                "onhand": stock.get(code, 0),
                "source": SOURCE,
            }
        )

    demands = []
    for rownum, (code, org, ship, qty) in enumerate(rows("报警表.xlsx"), 2):
        code = text(code)
        if code not in item_codes or num(qty) <= 0:
            continue
        demands.append(
            {
                "name": "FC_ALARM_%s_%s_%s_R%05d" % (org, code, str(ship).replace("/", ""), rownum),
                "customer": "FC_CUSTOMER",
                "item": code,
                "location": LOCATION,
                "due": iso_datetime(ship),
                "status": "open",
                "quantity": num(qty),
                "priority": 10,
                "source": SOURCE,
            }
        )
    for rownum, row in enumerate(unin_rows, 2):
        code = row["code"]
        if code not in item_codes or row["qty"] <= 0:
            continue
        order = row["order"] or "NA"
        due = iso_date_only(row["end"], days_offset=1) or "2026-09-18T00:00:00"
        demands.append(
            {
                "name": "FC_UNIN_%s_%s_%s_R%05d" % (row["org"] or "NA", code, order, rownum),
                "description": "未入库待排 %s %s" % (order, row["machine"] or row["workshop"] or ""),
                "customer": "FC_CUSTOMER",
                "item": code,
                "location": LOCATION,
                "due": due,
                "status": "open",
                "quantity": row["qty"],
                "priority": 1,
                "source": SOURCE,
            }
        )
    for rownum, (code, org, w1, w2, w3, total) in enumerate(rows("备库需求.xlsx"), 2):
        code = text(code)
        if code not in item_codes or num(total) <= 0:
            continue
        demands.append(
            {
                "name": "FC_STOCK_%s_%s_R%05d" % (org, code, rownum),
                "customer": "FC_CUSTOMER",
                "item": code,
                "location": LOCATION,
                "due": "2026-09-30T00:00:00",
                "status": "closed",
                "quantity": num(total),
                "priority": 20,
                "source": SOURCE,
            }
        )

    items = [
        {"name": code, "description": item_name.get(code, code), "category": item_category.get(code), "source": SOURCE}
        for code in sorted(item_codes)
    ]
    return {
        "clean_all": True,
        "parameters": [
            {"name": "currentdate", "value": "2026-09-10 08:00:00"},
            {"name": "plan.solver", "value": "mrp"},
        ],
        "data": {
            "locations": [{"name": LOCATION, "source": SOURCE}, {"name": "FC_RAW", "source": SOURCE}],
            "calendars": calendars,
            "calendarBuckets": buckets,
            "items": items,
            "resources": resources,
            "operations": operations,
            "operationResources": operation_resources,
            "operationMaterials": operation_materials,
            "buffers": buffers,
            "customers": [{"name": "FC_CUSTOMER", "source": SOURCE}],
            "demands": demands,
            "skills": [],
            "resourceSkills": [],
            "itemDistributions": [],
        },
        "_meta": {
            "bom_make_stocked": len(bom_make),
            "bom_components_stocked": len(bom_components),
            "sync_groups": len(sync_groups),
            "unin_rows": len(unin_rows),
            "unin_qty": sum(r["qty"] for r in unin_rows),
            "no_equipment": [
                {"code": code, "name": item_name.get(code, code), "workshop": item_workshop.get(code), "has_mold": bool(mold_map.get(code)), "has_ct": code in duration_per}
                for code in no_equipment
            ]
        },
    }


def main():
    doc = build()
    meta = doc.pop("_meta", {})
    out = ROOT / "fc_aps_scenario6.json"
    out.write_text(json.dumps(doc, ensure_ascii=False), encoding="utf-8")
    data = doc["data"]
    print(
        "items", len(data["items"]),
        "resources", len(data["resources"]),
        "operations", len(data["operations"]),
        "operationResources", len(data["operationResources"]),
        "operationMaterials", len(data["operationMaterials"]),
        "demands", len(data["demands"]),
        "calendars", len(data["calendars"]),
        "bom_make_stocked", meta.get("bom_make_stocked", 0),
        "bom_components_stocked", meta.get("bom_components_stocked", 0),
        "sync_groups", meta.get("sync_groups", 0),
        "unin_rows", meta.get("unin_rows", 0),
        "unin_qty", meta.get("unin_qty", 0),
    )
    missing = meta.get("no_equipment", [])
    print("no_equipment_on_order_items", len(missing))
    for row in missing:
        print(
            "  ",
            row["code"],
            row["name"],
            "工段=" + (row["workshop"] or "未知"),
            "模具=" + ("有" if row["has_mold"] else "无"),
            "CT=" + ("有" if row["has_ct"] else "无"),
        )
    csv_path = ROOT / "no_equipment.csv"
    with csv_path.open("w", encoding="utf-8") as f:
        f.write("物料编码,名称,工段,模具,CT\n")
        for row in missing:
            f.write(
                "%s,%s,%s,%s,%s\n"
                % (
                    row["code"],
                    row["name"].replace(",", "，"),
                    row["workshop"] or "未知",
                    "有" if row["has_mold"] else "无",
                    "有" if row["has_ct"] else "无",
                )
            )


if __name__ == "__main__":
    main()
