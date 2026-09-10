#!/usr/bin/env python3
"""Build the battery-factory CCPL request from the 2D-line Excel workbook.

Conversion SOP
==============
Input workbook sheets:
  * 物料表       -> materials: material code/name/type/cost/location/on-hand/UOM.
  * 工艺工序     -> operations and operationMaterials. One operation may occupy
                   multiple rows: rows are grouped by 工序ID and each material
                   row becomes one consume/produce flow.
  * 工作排期     -> resources and daily resource calendars. A 1 means an
                   08:00-18:00 working day; 0 means non-working. Resources
                   referenced by a process but absent from this sheet are
                   added as always-available fallback resources.
  * 订单需求表   -> open demands, due at 18:00 on the requested date.

Production modelling rules applied after extraction:
  * Workbook minutes are converted to CCPL seconds.
  * OP-01 through OP-11 model a continuous line: intermediate WIP producer
    and consumer flows have issueTiming=transfer_batch and transferBatch=1.
  * OP-02 workbook quantity 50 describes throughput, not a 50:1 conversion;
    its material flows stay 1:1 and it has no maximum batch limit.
  * OP-12 is the final discrete oven: sizeMinimum=sizeMultiple=sizeMaximum=128
    and its material flows are normalized to 1:1.
  * The resulting request uses constrained MRP (capa,mfg_lt,po_lt) for
    scenario5. It is data only: submitting it remains an MCP responsibility.

Usage:
  venv/bin/python scripts/build_battary_factory_request.py \
    --input 电池工厂二维产线数据清单.xlsx --output battary_factory.json
"""

import argparse
import json
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


DATASET_ID = "ds_battery_2d_line_001"
TENANT_ID = "battery"
LOCATION = "电池工厂二维产线"
STAMP = "2026-06-25 00:00:00"


def slug(value):
    """Produce stable CCPL IDs while retaining human-readable codes separately."""
    return re.sub(r"[^a-z0-9]+", "_", str(value).lower()).strip("_")


def value(row, header, name, default=None):
    index = header.get(name)
    return row[index] if index is not None and index < len(row) and row[index] is not None else default


def rows(sheet):
    header = {str(cell.value).strip(): index for index, cell in enumerate(sheet[1]) if cell.value is not None}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if any(cell is not None for cell in row):
            yield header, row


def as_date(value):
    return value.date() if isinstance(value, datetime) else value


def material_type(source):
    return {"原材料": "RAW", "中间在制品": "WIP", "成品": "FG"}.get(str(source), "WIP")


def build(workbook):
    now = {"createdUserId": "xlsx_import", "datasetId": DATASET_ID, "gmtCreated": STAMP, "gmtUpdated": STAMP, "tenantId": TENANT_ID}
    sheets = {sheet.title: sheet for sheet in workbook.worksheets}
    required = {"物料表", "工艺工序", "工作排期", "订单需求表"}
    missing = required.difference(sheets)
    if missing:
        raise ValueError("Workbook is missing sheets: %s" % ", ".join(sorted(missing)))

    materials, material_ids = [], {}
    for header, row in rows(sheets["物料表"]):
        code = str(value(row, header, "物料ID")).strip()
        material_id = "mat_%s" % slug(code)
        material_ids[code] = material_id
        material_ids[str(value(row, header, "物料名称")).strip()] = material_id
        materials.append({**now, "id": material_id, "materialCode": code, "materialName": value(row, header, "物料名称"), "materialType": material_type(value(row, header, "物料类型")), "unitCost": float(value(row, header, "单位成本", 0)), "locationName": value(row, header, "存放地点", LOCATION), "currentStock": float(value(row, header, "当前库存", 0)), "unitName": value(row, header, "计量单位", "个"), "remark": value(row, header, "物料类型", "")})

    resources, resource_ids, resource_calendars = [], {}, []
    for header, row in rows(sheets["工作排期"]):
        code = str(value(row, header, "识别符")).strip()
        resource_id = "res_%s" % slug(code)
        resource_ids[code] = resource_id
        name = value(row, header, "资源名称")
        skill = value(row, header, "技能", "通用") or "通用"
        resources.append({**now, "id": resource_id, "resourceCode": code, "resourceName": name, "bindDeviceName": name, "resourceType": "default", "skillName": skill, "locationName": LOCATION, "capacity": 1.0, "efficiency": 1.0, "isEnabled": 1, "remark": "由工作排期表导入"})
        for date_name, index in header.items():
            if isinstance(sheets["工作排期"][1][index].value, datetime):
                day = as_date(sheets["工作排期"][1][index].value)
                resource_calendars.append({**now, "id": "cal_%s_%s" % (slug(code), day.strftime("%Y%m%d")), "resourceId": resource_id, "bindDeviceName": name, "calendarDate": day.isoformat(), "isWorkingDay": int(bool(row[index])), "shiftStartTime": "08:00:00", "shiftEndTime": "18:00:00", "workCapacity": 1.0, "skillName": skill, "remark": "由工作排期表导入"})

    operations, operation_materials, operation_by_code = [], [], {}
    sequence = 0
    grouped = defaultdict(list)
    for header, row in rows(sheets["工艺工序"]):
        grouped[str(value(row, header, "工序ID")).strip()].append((header, row))
    for code, group in grouped.items():
        sequence += 10
        header, first = group[0]
        resource_name = value(first, header, "资源名称")
        # The work calendar identifies resources by code, while the route names
        # them by display name. Create a deterministic fallback when no code matches.
        resource_id = next((r["id"] for r in resources if r["resourceName"] == resource_name), None)
        if resource_id is None:
            resource_id = "res_%s" % slug(resource_name)
            if resource_id not in {r["id"] for r in resources}:
                resources.append({**now, "id": resource_id, "resourceCode": resource_name, "resourceName": resource_name, "bindDeviceName": resource_name, "resourceType": "default", "skillName": value(first, header, "所需技能", "通用") or "通用", "locationName": LOCATION, "capacity": 1.0, "efficiency": 1.0, "isEnabled": 1, "remark": "工艺表补充资源"})
        operation_id = "op_%s" % slug(code)
        operation_by_code[code] = operation_id
        operations.append({**now, "id": operation_id, "operationCode": code, "operationName": value(first, header, "工序名称"), "operationType": value(first, header, "工序类型"), "fixedDurationSeconds": int(float(value(first, header, "固定生产时间（min）", 0)) * 60), "unitDurationSeconds": int(float(value(first, header, "每件生产时间（min）", 0)) * 60), "priority": int(value(first, header, "优先级", sequence // 10)), "sequenceNo": sequence, "resourceId": resource_id, "resourceQuantity": 1.0, "requiredSkillName": value(first, header, "所需技能", "通用") or "通用", "routeCode": "ROUTE_BATTERY_2D_LINE_001", "routeName": "电池工厂二维产线工艺路线", "isEnabled": 1, "remark": value(first, header, "工序类型")})
        for index, (header, row) in enumerate(group, 1):
            item_code = str(value(row, header, "物料ID", "")).strip()
            item_name = str(value(row, header, "物料名称")).strip()
            material_key = item_code or item_name
            material_id = material_ids.get(material_key)
            if material_id is None:
                material_id = "mat_%s" % slug(material_key)
                material_ids[material_key] = material_id
                materials.append({**now, "id": material_id, "materialCode": item_code or item_name, "materialName": item_name, "materialType": material_type(value(row, header, "类型")), "unitCost": 0.0, "locationName": LOCATION, "currentStock": 0.0, "unitName": value(row, header, "单位", "个"), "remark": "由工艺表生成的在制品"})
            produce = value(row, header, "产出或消耗") == "产出"
            operation_materials.append({**now, "id": "opm_%s_%03d" % (slug(code), index), "operationId": operation_id, "materialId": material_id, "materialRole": "produce" if produce else "consume", "issueTiming": "end" if produce else "start", "quantity": float(value(row, header, "消耗与生成量", 1)), "priority": 1, "unitName": value(row, header, "单位", "个"), "remark": "%s %s %s" % (value(row, header, "工序名称"), "产出" if produce else "消耗", item_name)})

    demands = []
    for header, row in rows(sheets["订单需求表"]):
        code = str(value(row, header, "需求标识")).strip()
        due = value(row, header, "截止日期")
        due_text = (as_date(due).isoformat() if isinstance(due, datetime) else str(due)) + " 18:00:00"
        material_code = str(value(row, header, "需求物料ID")).strip()
        demands.append({**now, "id": "dem_%s" % slug(code), "demandCode": code, "materialId": material_ids[material_code], "customerName": value(row, header, "客户名称"), "deliveryLocationName": "暂存仓2", "dueTime": due_text, "quantity": float(value(row, header, "需求数量")), "priority": int(value(row, header, "需求优先级", 10)), "status": "OPEN", "remark": "由订单需求表导入"})

    # Apply the line-specific transfer and batch rules documented in this file.
    by_material = {flow["materialId"]: flow for flow in operation_materials if flow["materialRole"] == "produce"}
    by_id = {op["id"]: op for op in operations}
    for flow in operation_materials:
        producer = by_material.get(flow["materialId"])
        if flow["materialRole"] == "consume" and producer:
            upstream, downstream = by_id[producer["operationId"]], by_id[flow["operationId"]]
            if upstream["operationCode"] != "OP-12" and downstream["operationCode"] != "OP-12":
                producer.update(issueTiming="transfer_batch", transferBatch=1)
                flow.update(issueTiming="transfer_batch", transferBatch=1)
    for op in operations:
        if op["operationCode"] == "OP-12":
            op.update(sizeMinimum=128, sizeMultiple=128, sizeMaximum=128)
            for flow in operation_materials:
                if flow["operationId"] == op["id"]:
                    flow["quantity"] = 1.0
        elif op["operationCode"] == "OP-02":
            op.pop("sizeMaximum", None)
            for flow in operation_materials:
                if flow["operationId"] == op["id"]:
                    flow["quantity"] = 1.0

    data = {"dataset": {**now, "id": DATASET_ID, "code": "BATTERY_2D_LINE_001", "name": "电池工厂二维产线数据清单", "status": 1, "versionNo": 1, "remark": "由电池工厂二维产线数据清单.xlsx转换"}, "defaultLocationName": LOCATION, "materials": materials, "resources": resources, "resourceCalendars": resource_calendars, "calendarsByResourceId": {r["id"]: [c for c in resource_calendars if c["resourceId"] == r["id"]] for r in resources}, "operations": operations, "operationMaterials": operation_materials, "operationMaterialsByOperationId": {op["id"]: [f for f in operation_materials if f["operationId"] == op["id"]] for op in operations}, "demands": demands}
    return {"request_id": "battery_2d_line_request", "scenario": "scenario5", "clean": True, "clean_all": True, "constraint": "capa,mfg_lt,po_lt", "plantype": 1, "env": "supply", "data": data}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", default="电池工厂二维产线数据清单.xlsx")
    parser.add_argument("--output", default="battary_factory.json")
    args = parser.parse_args()
    request = build(load_workbook(args.input, data_only=True))
    Path(args.output).write_text(json.dumps(request, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    data = request["data"]
    print("wrote %s: %s materials, %s resources, %s operations, %s flows, %s demands" % (args.output, len(data["materials"]), len(data["resources"]), len(data["operations"]), len(data["operationMaterials"]), len(data["demands"])))


if __name__ == "__main__":
    main()
