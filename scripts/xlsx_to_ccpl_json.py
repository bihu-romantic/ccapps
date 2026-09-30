#!/usr/bin/env python3
"""Convert a ccAPPS 13-table workbook (.xlsx) into a ccpl JSON payload.

The JSON keeps the Excel sheet layout so ``/api/ccpl/plan/`` can rebuild
the workbook and reuse the importworkbook engine.

Usage:
  python scripts/xlsx_to_ccpl_json.py \\
      /mnt/c/Users/CC/Desktop/fixed_data_13tables.xlsx \\
      /home/c/ccapps/fixed_data_13tables.json
"""

from __future__ import annotations

import argparse
import json
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

# Chinese Excel tab → English sheet key (ccpl / standard-dataset style).
# Keys must be accepted by matchesModelName (or remapped at import).
SHEET_NAME_MAP = {
    "地点": "locations",
    "日历": "calendars",
    "calendar": "calendars",
    "日历-排班表": "calendarBuckets",
    "日历：排班表": "calendarBuckets",
    "客户": "customers",
    "对象": "items",
    "库存": "buffers",
    "技能": "skills",
    "资源": "resources",
    "资源技能": "resourceSkills",
    "操作": "operations",
    "操作-物料": "operationMaterials",
    "操作：物料": "operationMaterials",
    "操作-资源": "operationResources",
    "操作：资源": "operationResources",
    "销售订单": "demands",
}


def _json_value(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float, str)):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, timedelta):
        total = int(value.total_seconds())
        sign = "-" if total < 0 else ""
        total = abs(total)
        h, rem = divmod(total, 3600)
        m, s = divmod(rem, 60)
        return "%s%02d:%02d:%02d" % (sign, h, m, s)
    return str(value)


def workbook_to_sheets(path: Path) -> dict:
    wb = load_workbook(filename=str(path), data_only=True)
    sheets = {}
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(values_only=True))
        if not rows:
            key = SHEET_NAME_MAP.get(name, name)
            sheets[key] = []
            continue
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        records = []
        for row in rows[1:]:
            if row is None or all(c is None or c == "" for c in row):
                continue
            record = {}
            for i, header in enumerate(headers):
                if not header:
                    continue
                record[header] = _json_value(row[i] if i < len(row) else None)
            records.append(record)
        key = SHEET_NAME_MAP.get(name, name)
        sheets[key] = records
    return sheets


def _stamp_source(sheets: dict, source: str) -> dict:
    """Ensure every row's 源 matches the envelope source (replace_source)."""
    stamped = {}
    for name, rows in sheets.items():
        out = []
        for row in rows:
            item = dict(row)
            item["源"] = source
            out.append(item)
        stamped[name] = out
    return stamped


def build_payload(
    xlsx_path: Path,
    scenario: str = "scenario2",
    *,
    source: str | None = None,
    currentdate: str = "2026-07-15 08:00:00",
    clean_all: bool = True,
    solver: str = "ga",
) -> dict:
    """Build a ccpl protocol payload: envelope + 13 Excel sheets."""
    dataset_id = source or ("%s-13tables" % scenario)
    source = dataset_id
    sheets = _stamp_source(workbook_to_sheets(xlsx_path), source)
    return {
        "version": "1.0",
        "requestId": "aps-%s-%s-0001"
        % (scenario, datetime.now().strftime("%Y%m%d%H%M")),
        "datasetId": dataset_id,
        "scenario": scenario,
        "source": source,
        "mode": "replace_source",
        "clean_all": clean_all,
        "parameters": [
            {
                "name": "currentdate",
                "value": currentdate,
            }
        ],
        "plan_options": {
            "constraint": "capa,mfg_lt,po_lt",
            "plantype": 1,
            "env": "supply",
            "solver": {
                "type": solver,
                "objectiveProfile": "due_date",
            },
            "objectives": [
                {"code": "DELIVERY_ACCURACY", "weight": 73},
                {"code": "EQUIPMENT_LOAD_BALANCE", "weight": 15},
                {"code": "MINIMIZE_CHANGEOVER", "weight": 12},
            ],
        },
        "sheets": sheets,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("json_out", type=Path)
    parser.add_argument("--scenario", default="scenario2")
    parser.add_argument("--source", default=None)
    parser.add_argument("--currentdate", default="2026-07-15 08:00:00")
    parser.add_argument("--solver", default="ga")
    parser.add_argument(
        "--no-clean-all",
        action="store_true",
        help="Keep existing data outside this source (default: clean_all=true)",
    )
    parser.add_argument(
        "--pretty",
        action="store_true",
        help="Pretty-print JSON (default is compact for smaller uploads)",
    )
    args = parser.parse_args()

    payload = build_payload(
        args.xlsx,
        scenario=args.scenario,
        source=args.source,
        currentdate=args.currentdate,
        clean_all=not args.no_clean_all,
        solver=args.solver,
    )
    args.json_out.parent.mkdir(parents=True, exist_ok=True)
    if args.pretty:
        text = json.dumps(payload, ensure_ascii=False, indent=2)
    else:
        text = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    args.json_out.write_text(text, encoding="utf-8")
    sheets = payload["sheets"]
    print(
        "Wrote %s (%s, %d sheets, %d rows)"
        % (
            args.json_out,
            sizeof_fmt := "%0.1f MB" % (args.json_out.stat().st_size / 1024 / 1024),
            len(sheets),
            sum(len(v) for v in sheets.values()),
        )
    )
    for name, rows in sheets.items():
        print("  %-12s %d" % (name, len(rows)))


if __name__ == "__main__":
    main()
